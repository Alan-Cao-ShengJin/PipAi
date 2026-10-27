"""Insert a 4Q26E forecast column in 3S (after 3Q26) and make FY2026E = 9M26 actual + 4Q26E."""
import json, os, sys
from copy import copy
import openpyxl
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as gl, column_index_from_string as ci

MODEL = r"C:\Users\Alan Cao\Desktop\Companies\宇通\宇通客车 1015.xlsx"
SCR = os.path.dirname(os.path.abspath(__file__))
R = json.load(open(os.path.join(SCR, "rp_rows.json")))
RP_Q4 = gl(11 + 11)          # Revenue Projection column of 4Q26E (quarters start at K=11; 4Q26 is index 11)
wb = openpyxl.load_workbook(MODEL)
ws = wb["3S"]
LAST = 58
PURPLE = "FF7030A0"
assert ws["S1"].value is None and ws["T1"].value == "1H25", (ws["S1"].value, ws["T1"].value)

# 1) shift T..X -> U..Y (values, styles, comments)
for col in ["X", "W", "V", "U", "T"]:
    new = gl(ci(col) + 1)
    for r in range(1, LAST + 1):
        a, b = ws[f"{col}{r}"], ws[f"{new}{r}"]
        v = a.value
        b.value = Translator(v, origin=f"{col}{r}").translate_formula(f"{new}{r}") if isinstance(v, str) and v.startswith("=") else v
        b.font, b.fill, b.border, b.alignment, b.number_format = copy(a.font), copy(a.fill), copy(a.border), copy(a.alignment), a.number_format
        b.comment = copy(a.comment) if a.comment else None
        a.value, a.comment = None, None
for c in "SUVXY":
    ws.column_dimensions[c].width = 13
for c in "TW":
    ws.column_dimensions[c].width = 3
# blank gap column T styles
for r in range(1, LAST + 1):
    t = ws[f"T{r}"]
    t.font, t.fill, t.border = Font(), openpyxl.styles.PatternFill(), openpyxl.styles.Border()

# 2) 4Q26E column S: styles from R (3Q26), header from H (EST预测)
for r in range(1, LAST + 1):
    a, b = ws[f"R{r}"], ws[f"S{r}"]
    b.font, b.fill, b.border, b.alignment, b.number_format = copy(a.font), copy(a.fill), copy(a.border), copy(a.alignment), a.number_format
ws["S1"] = "4Q26E"
ws["S2"] = "EST预测"
ws["S2"].font = copy(ws["H2"].font)
for r in range(3, LAST + 1):
    v = ws[f"R{r}"].value
    if isinstance(v, str) and v.startswith("="):
        ws[f"S{r}"] = Translator(v, origin=f"R{r}").translate_formula(f"S{r}")

ASM = {}   # cell -> rationale
def put(ref, val, why=None):
    ws[ref] = val
    if why:
        ASM[ref] = why
put("S4", f"='Revenue Projection'!{RP_Q4}{R['rev_tot']}*100000")
put("S6", f"=-'Revenue Projection'!{RP_Q4}{R['cost_tot']}*100000")
for amt, ratio, lab in ((10, 11, "税金及附加"), (12, 13, "销售费用"), (14, 15, "管理费用"), (16, 17, "研发费用")):
    put(f"S{amt}", f"=-S{ratio}*S4")
    put(f"S{ratio}", f"=O{ratio}", f"4Q26E {lab}率=4Q25实际（Q4费用季节性偏高）")
put("S18", "=Y18/3", "其他收益=9M26季度均值（政府补助，无明显季节性）")
put("S20", "=Y20/3", "其他投资收益=9M26季度均值（含应收款贴现损失）")
put("S22", 0, "公允价值变动无法预测，取0")
put("S24", "=O24", "信用减值=4Q25实际（年末集中计提）")
put("S26", "=O26", "资产减值=4Q25实际（年末集中计提）")
put("S28", 0, "资产处置取0")
put("S30", "=O30", "营业外净收入=4Q25实际")
put("S40", "=Y40/3", "财务费用=9M26季度均值")
put("S41", "=Y41/3", "利息收入=9M26季度均值")
put("S44", "=Y44/3", "联营/合营投资收益=9M26季度均值")
put("S47", "=-H46*H48-Y47")
put("S48", "=-S47/S46")
put("S51", "=S50-S54")
put("S54", "=S55*S50")
put("S55", "=G55", "少数股东损益占比=2025全年")
put("S5", "=S4/O4-1")
put("S52", "=S51/O51-1")

# 3) FY2026E (H) = 9M26 (Y) + 4Q26E (S) for flow lines; ratios become outputs
for r in (10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 30, 40, 41, 42, 44, 54):
    ws[f"H{r}"] = f"=Y{r}+S{r}"
for r, f in ((11, "=-H10/H4"), (13, "=-H12/H4"), (15, "=-H14/H4"), (17, "=-H16/H4"), (21, "=-H20/H4"), (55, "=H54/H50")):
    ws[f"H{r}"] = f
put("H48", "=G48", "全年所得税率与2025年持平（Guidance #9：全年税率和去年差不多）；4Q26E税率倒算")
# 2027-28E: non-recurring lines set to zero (were carrying 2026 one-offs forward)
for c in "IJ":
    put(f"{c}23", 0, "公允价值变动不可预测，取0")
    put(f"{c}25", 0, "信用减值：2025-26为冲回（非经常），预测取0")
    put(f"{c}29", 0, "资产处置取0")

# fonts: assumptions purple, outputs black (keep size/name)
def recolor(ref, purple):
    f = copy(ws[ref].font)
    ws[ref].font = Font(name=f.name, size=f.sz, bold=f.b, italic=f.i, color=PURPLE if purple else None)
for ref in ASM:
    recolor(ref, True)
    ws[ref].comment = Comment(ASM[ref], "Claude")
for ref in ["H11", "H13", "H15", "H17", "H18", "H21", "H44", "H55", "H10", "H12", "H14", "H16", "H20", "H22", "H24", "H26", "H28", "H30", "H40", "H54"]:
    if ref not in ASM:
        recolor(ref, False)
ws["S1"].comment = Comment("4Q26E预测：收入/成本取自Revenue Projection 4Q26E；费用率取4Q25；其他项目取9M26季度均值或4Q25；所得税按全年税率倒算。FY2026E=9M26实际+4Q26E。", "Claude")

wb.calculation.fullCalcOnLoad = True
wb.save(MODEL)
print("done; RP 4Q26 col", RP_Q4, "| assumptions:", len(ASM))
