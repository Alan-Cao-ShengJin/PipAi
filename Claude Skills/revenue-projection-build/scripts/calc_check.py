"""Calculate a copy of an Excel model and check it.

Usage:
  python calc_check.py <model.xlsx> [--sheet "Revenue Projection"] [--rows 4-60] [--cols C:I,K:Z]
                       [--backup <backup.xlsx>] [--workdir <dir>]

- Copies the workbook to --workdir (default: system temp), calculates every formula with the
  `formulas` engine (pip install formulas), and reports any #DIV/0!, #REF!, #VALUE! ... per sheet.
- With --sheet/--rows/--cols, prints the calculated values of that block (label from column A or B).
- With --backup, lists the cells whose formulas/values differ from the backup, per sheet
  (use it to prove that only the intended tabs changed).
openpyxl does not store calculated values, so always use this rather than reading data_only from
a file openpyxl just wrote. Excel will recalculate on open if fullCalcOnLoad is set.
"""
import argparse, os, shutil, subprocess, sys, tempfile

try:
    import formulas
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "formulas", "-q"])
    import formulas
import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl


def col_list(spec):
    out = []
    for part in spec.split(","):
        a, _, b = part.partition(":")
        out += [gl(c) for c in range(ci(a), ci(b or a) + 1)]
    return out


def main():
    p = argparse.ArgumentParser()
    p.add_argument("model")
    p.add_argument("--sheet")
    p.add_argument("--rows", default="1-80")
    p.add_argument("--cols", default="C:I")
    p.add_argument("--backup")
    p.add_argument("--workdir", default=tempfile.gettempdir())
    a = p.parse_args()

    os.makedirs(a.workdir, exist_ok=True)
    src = os.path.join(a.workdir, "calc_in.xlsx")
    out = os.path.join(a.workdir, "calc_out")
    shutil.copy2(a.model, src)
    shutil.rmtree(out, ignore_errors=True)
    xl = formulas.ExcelModel().loads(src).finish(circular=True)
    xl.calculate()
    xl.write(dirpath=out)
    calc = openpyxl.load_workbook(os.path.join(out, os.path.basename(src).upper()), data_only=True)
    raw = openpyxl.load_workbook(src)

    for s in raw.sheetnames:
        ws = calc[s.upper()]
        codes = {"#DIV/0!", "#REF!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
        errs = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.strip().upper() in codes]
        print(f"[{s}] errors: {len(errs)} {errs[:10]}")

    if a.sheet:
        ws, wf = calc[a.sheet.upper()], raw[a.sheet]
        r0, r1 = (int(x) for x in a.rows.split("-"))
        cols = col_list(a.cols)
        print("row label".ljust(34), *[str(wf[f"{c}1"].value or c)[:8].rjust(9) for c in cols])
        for r in range(r0, r1 + 1):
            lab = wf.cell(r, 1).value or wf.cell(r, 2).value
            if lab is None:
                continue
            vals = []
            for c in cols:
                v, fmt = ws[f"{c}{r}"].value, wf[f"{c}{r}"].number_format
                if isinstance(v, (int, float)):
                    vals.append(f"{v:9.1%}" if "%" in fmt else f"{v:9.0f}" if abs(v) >= 1000 else f"{v:9.2f}")
                else:
                    vals.append(" " * 9)
            print(f"{r:4} {str(lab).strip()[:28]:28}", *vals)

    if a.backup:
        b = openpyxl.load_workbook(a.backup)
        for s in b.sheetnames:
            if s not in raw.sheetnames:
                print(f"[{s}] missing now"); continue
            d = [c.coordinate for row in b[s].iter_rows() for c in row if c.value != raw[s][c.coordinate].value]
            print(f"[{s}] cells changed vs backup: {len(d)} {d[:15]}")
        for s in raw.sheetnames:
            if s not in b.sheetnames:
                print(f"[{s}] new sheet")


if __name__ == "__main__":
    main()
