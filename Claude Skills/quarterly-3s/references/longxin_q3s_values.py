"""Compute 3S quarterly / 1H / 9M input rows (RMB'000) for 隆鑫通用 from the 2026 Q1/H1/Q3 reports + FY25 in the model."""
import json, os, sys
import openpyxl
S=sys.argv[1]; F=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
J={k:json.load(open(os.path.join(S,f'is_{k}.json'),encoding='utf-8')) for k in ['q1','h1','q3']}
def line(d,key):
    for k,v in d['income_statement'].items():
        if k.endswith(key): return v
    raise KeyError(key)
def per(k,which):
    d=J[k]; g=lambda key: line(d,key)[which]/1000
    other=[g('其他收益'),g('投资收益'),g('信用减值损失'),g('资产减值损失'),g('资产处置收益')]
    return {4:g('营业收入'),6:-g('营业成本'),10:-g('税金及附加'),12:-g('销售费用'),14:-g('管理费用'),16:-g('研发费用'),
            18:-g('财务费用'),20:sum(other),22:g('营业外收入')-g('营业外支出'),33:-g('所得税费用'),
            37:g('归属于母公司股东的净利润'),'_other':other,'_nonop':[g('营业外收入'),g('营业外支出')],
            '_pbt':g('利润总额'),'_jv':g('对联营企业和合营企业的投资收益')}
P={'1Q26':per('q1',0),'1Q25':per('q1',1),'1H26':per('h1',0),'1H25':per('h1',1),'9M26':per('q3',0),'9M25':per('q3',1)}
# FY25 from the model (annual report), column G
wb=openpyxl.load_workbook(os.path.join(F,'隆鑫通用 0630.xlsx'),data_only=True); ws=wb['3S']
P['FY25']={r:ws[f'G{r}'].value for r in [4,6,10,12,14,16,18,20,22,33,37]}
P['FY25']['_pbt']=ws['G32'].value
R=[4,6,10,12,14,16,18,20,22,33,37]
def diff(a,b):
    out={r:P[a][r]-P[b][r] for r in R}; out['_pbt']=P[a]['_pbt']-P[b]['_pbt']
    out['_jv']=(P[a].get('_jv',0) or 0)-(P[b].get('_jv',0) or 0) if '_jv' in P[a] else None
    return out
Q={'1Q25':P['1Q25'],'2Q25':diff('1H25','1Q25'),'3Q25':diff('9M25','1H25'),'4Q25':diff('FY25','9M25'),
   '1Q26':P['1Q26'],'2Q26':diff('1H26','1Q26'),'3Q26':diff('9M26','1H26')}
# check 1Q columns already in model
print('model L/M vs report 1Q:',[round(ws[f'L{r}'].value-P['1Q25'][r],3) for r in R],[round(ws[f'M{r}'].value-P['1Q26'][r],3) for r in R])
# PBT identity per period: GP + rows = 利润总额
for k,v in {**Q,'1H25':P['1H25'],'1H26':P['1H26'],'9M25':P['9M25'],'9M26':P['9M26']}.items():
    pbt=sum(v[r] for r in [4,6,10,12,14,16,18,20,22]); print(f'{k}: PBT build {pbt:,.1f} vs report {v["_pbt"]:,.1f} diff {pbt-v["_pbt"]:.3f}')
FS={'rev':[4558.60946484,5105.63718434,4804.99449701,4578.20699787,4522.34354416,5252.89915097,5086.26863599],
    'ni':[506.88111402,567.33974567,502.58912846,70.94038261,471.33495238,519.87760515,551.79999934],
    'tax':[90.69171382,100.82739198,94.96066609,8.83010998,89.8561537,95.44430802,97.20461877],
    'ptx':[593.2029817,666.58391163,595.76284496,72.32276482,566.27741452,618.05148154,650.82319153],
    'nontax':[0,112.04541037,113.02431916,96.31458556,99.10960121,141.41154949,129.83604828]}
print('\nFactSet check (model − FactSet, RMB m):')
for i,k in enumerate(Q):
    v=Q[k]; jv=v['_jv']
    print(k, 'rev',round(v[4]/1000-FS['rev'][i],3),'ni',round(v[37]/1000-FS['ni'][i],3),'tax',round(-v[33]/1000-FS['tax'][i],3),
          'nontax',round(-v[10]/1000-FS['nontax'][i],3),'pbt−jv',round((v['_pbt']-(jv or 0))/1000-FS['ptx'][i],3) if jv is not None else 'n/a (FY JV n/a)', f'(PBT {v["_pbt"]/1000:.1f})')
out={'quarters':Q,'cum':{k:P[k] for k in ['1H25','1H26','9M25','9M26']}}
json.dump(out,open(os.path.join(S,'lx_q3s_values.json'),'w',encoding='utf-8'),ensure_ascii=False,indent=1)
print('\nsummary  rev(亿) GPM  ETR  NI(亿)')
for k,v in {**Q,'1H26':P['1H26'],'9M26':P['9M26']}.items():
    print(f'{k:5s} {v[4]/1e5:6.2f} {(v[4]+v[6])/v[4]:.1%} {-v[33]/v["_pbt"]:.1%} {v[37]/1e5:5.2f}')
