"""Fill the interim (quarterly / 1H / 9M) columns of a 3S income statement from cumulative report data.

Usage:
  python fill_quarters.py <config.json> [--dry-run]

See SKILL.md for the config format. In short:
  - "cumulative" periods come either from a column already in the sheet ({"col": "M"}) or from an
    extract_is.py JSON ({"json": "...", "which": 0|1}); 1Q periods are cumulative = quarterly.
  - "quarters" are single-quarter periods written as HARD VALUES = period A - period B
    (e.g. "2Q25": ["1H25", "1Q25"], "4Q25": ["FY25", "9M25"]).
  - "fy" maps annual labels to their existing column (e.g. {"FY25": "G"}).
  - "layout" lists labels left to right starting at "first_col"; "|" is a narrow gap column.
Ratio rows (margins, % of revenue, ETR, PBT, PAT...) are copied as formulas from "template_col".
Formulas elsewhere in the sheet that pointed at a moved interim column are repointed.
Safety: refuses if the workbook is open (~$ lock file), backs up first, prints a reconciliation.
"""
import json
import os
import re
import shutil
import sys
from copy import copy
from datetime import datetime

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl

LABELS = {  # model row key -> label prefixes searched in column B (case-insensitive)
    "rev": ["revenue"], "cogs": ["operating costs", "cost of sales"], "tax_s": ["business tax"],
    "sell": ["selling expenses"], "mgmt": ["management expenses", "administrative expenses"],
    "rd": ["r&d expenses"], "oth": ["other income"], "inv": ["other investment gains", "investment income"],
    "fv": ["net income from fair value", "fair value"], "credit": ["credit imp"], "asset": ["assets impairment", "asset impairment"],
    "disp": ["asset disposal"], "nonop": ["non-operating"], "fin": ["finance expense"], "int": ["finance income"],
    "jv": ["share of profit"], "itax": ["income tax expense"], "ni": ["equity shareholders"],
}
IS_LABELS = {  # model key -> 合并利润表 label(s)
    "rev": ["营业收入"], "cogs": ["营业成本"], "tax_s": ["税金及附加"], "sell": ["销售费用"], "mgmt": ["管理费用"],
    "rd": ["研发费用"], "fin": ["财务费用"], "int": ["利息收入"], "oth": ["其他收益"], "inv": ["投资收益"],
    "jv": ["对联营企业和合营企业的投资收益"], "fv": ["公允价值变动收益"], "credit": ["信用减值损失"],
    "asset": ["资产减值损失"], "disp": ["资产处置收益"], "nop_in": ["营业外收入"], "nop_out": ["营业外支出"],
    "itax": ["所得税费用"], "ni": ["归属于母公司股东的净利润", "归属于母公司所有者的净利润"],
}
NEG = {"cogs", "tax_s", "sell", "mgmt", "rd", "fin", "itax"}  # model stores these as negatives


def find_rows(ws, overrides):
    rows = {}
    for r in range(1, ws.max_row + 1):
        lab = str(ws.cell(r, 2).value or "").strip().lower()
        for k, pre in LABELS.items():
            if k not in rows and any(lab.startswith(p) for p in pre):
                rows[k] = r
    rows.update({k: int(v) for k, v in overrides.items()})
    missing = [k for k in LABELS if k not in rows]
    if missing:
        sys.exit(f"Could not find rows for {missing}; add them to config 'rows'.")
    return rows


def is_values(path, which, rows, col):
    d = json.load(open(path, encoding="utf-8"))["income_statement"]
    def g(key):
        for lab in IS_LABELS[key]:
            if lab in d and d[lab][which] is not None:
                return round(d[lab][which] / 1000, 5)
        return 0.0
    out = {}
    for k in LABELS:
        if k == "inv":
            out[rows[k]] = f"={g('inv'):.5f}-{col}{rows['jv']}"
        elif k == "nonop":
            out[rows[k]] = f"={g('nop_in'):.5f}-{g('nop_out'):.5f}"
        else:
            v = g(k)
            out[rows[k]] = -v if k in NEG else v
    return out


class Evaluator:
    def __init__(self, ws):
        self.ws, self.cache = ws, {}

    def __call__(self, ref):
        if ref in self.cache:
            return self.cache[ref]
        v = self.ws[ref].value
        if isinstance(v, str) and v.startswith("="):
            f = v[1:]
            if "!" in f or re.search(r"[A-Z]{2,}\(", f.replace("SUM(", "")):
                v = None
            else:
                f = re.sub(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", lambda m: "(" + "+".join(
                    f"{gl(c)}{r}" for c in range(ci(m.group(1)), ci(m.group(3)) + 1)
                    for r in range(int(m.group(2)), int(m.group(4)) + 1)) + ")", f)
                f = re.sub(r"\$?([A-Z]{1,2})\$?(\d+)", lambda m: "(" + repr(self(m.group(1) + m.group(2)) or 0) + ")", f)
                f = f.replace("%", "/100")
                try:
                    v = eval(f)
                except ZeroDivisionError:
                    v = None
        self.cache[ref] = v
        return v


def main():
    cfg = json.load(open(sys.argv[1], encoding="utf-8"))
    dry = "--dry-run" in sys.argv
    model, sheet = cfg["model"], cfg.get("sheet", "3S")
    lock = os.path.join(os.path.dirname(model), "~$" + os.path.basename(model))
    if os.path.exists(lock) and not dry:
        sys.exit("ABORT: workbook is open in Excel. Close it and retry.")
    if not dry:
        bdir = cfg.get("backup_dir", os.path.dirname(model))
        stem, ext = os.path.splitext(os.path.basename(model))
        bk = os.path.join(bdir, f"{stem}.backup-{datetime.now():%Y%m%d-%H%M%S}{ext}")
        shutil.copy2(model, bk)
        print("backup:", bk)

    wb = openpyxl.load_workbook(model)
    ws = wb[sheet]
    rows = find_rows(ws, cfg.get("rows", {}))
    input_rows = sorted(rows.values())
    tcol = cfg["template_col"]
    lab_row, type_row = cfg.get("label_row", 1), cfg.get("type_row", 2)
    last_row = max(r for r in range(1, ws.max_row + 1) if ws[f"{tcol}{r}"].value is not None)
    tmpl = {r: ws[f"{tcol}{r}"].value for r in range(1, last_row + 1)}
    yoy_rows = {}
    for r, v in tmpl.items():
        m = re.fullmatch(rf"={tcol}(\d+)/[A-Z]+\1-1", str(v)) if isinstance(v, str) else None
        if m:
            yoy_rows[r] = int(m.group(1))
    formula_rows = [r for r, v in tmpl.items() if isinstance(v, str) and v.startswith("=")
                    and r not in input_rows and r not in yoy_rows and r > type_row]
    styles = {r: ws[f"{tcol}{r}"] for r in range(1, last_row + 1)}
    styles = {r: (copy(c.font), copy(c.fill), copy(c.border), copy(c.alignment), c.number_format) for r, c in styles.items()}

    # layout -> column letters
    layout, pos = cfg["layout"], {}
    c = ci(cfg["first_col"])
    for lab in layout:
        if lab != "|":
            pos[lab] = gl(c)
        c += 1
    last_col = c - 1
    area_cols = {gl(x) for x in range(ci(cfg["first_col"]), last_col + 1)}
    pos.update(cfg.get("fy", {}))

    # snapshot sources that live in the sheet
    snap = {lab: {r: ws[f"{s['col']}{r}"].value for r in input_rows}
            for lab, s in cfg["cumulative"].items() if "col" in s}
    moved = {s["col"]: pos[lab] for lab, s in cfg["cumulative"].items() if "col" in s and s["col"] != pos[lab]}

    # clear the interim area
    for r in range(1, last_row + 1):
        for col in area_cols:
            cell = ws[f"{col}{r}"]
            cell.value, cell.comment = None, None
            cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = Font(), PatternFill(), Border(), Alignment(), "General"

    def style(col):
        for r in range(1, last_row + 1):
            f, fl, b, al, nf = styles[r]
            cell = ws[f"{col}{r}"]
            cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = copy(f), copy(fl), copy(b), copy(al), nf

    # 1) cumulative periods
    for lab, s in cfg["cumulative"].items():
        col = pos[lab]
        style(col)
        if "col" in s:
            for r in input_rows:
                v = snap[lab][r]
                ws[f"{col}{r}"] = Translator(v, origin=f"{s['col']}{r}").translate_formula(f"{col}{r}") \
                    if isinstance(v, str) and v.startswith("=") else v
        else:
            for r, v in is_values(s["json"], s["which"], rows, col).items():
                ws[f"{col}{r}"] = v

    # 2) single quarters as hard values
    ev = Evaluator(ws)
    for lab, (a, b) in cfg.get("quarters", {}).items():
        col = pos[lab]
        style(col)
        for r in input_rows:
            va, vb = ev(f"{pos[a]}{r}"), ev(f"{pos[b]}{r}")
            ws[f"{col}{r}"] = round((va or 0) - (vb or 0), 5)

    # 3) labels, ratio formulas, YoY
    for lab, col in pos.items():
        if col not in area_cols:
            continue
        ws[f"{col}{lab_row}"] = lab
        ws[f"{col}{type_row}"] = cfg.get("type_label", "ACT实际")
        for r in formula_rows:
            ws[f"{col}{r}"] = Translator(tmpl[r], origin=f"{tcol}{r}").translate_formula(f"{col}{r}")
        base = cfg.get("yoy", {}).get(lab)
        for r, br in yoy_rows.items():
            ws[f"{col}{r}"] = f"={col}{br}/{pos[base]}{br}-1" if base else None
    for lab, text in cfg.get("comments", {}).items():
        ws[f"{pos[lab]}{lab_row}"].comment = Comment(text, "Claude")
    for col in area_cols:
        ws.column_dimensions[col].width = 13
    for x in range(ci(cfg["first_col"]), last_col + 1):
        if gl(x) not in pos.values():
            ws.column_dimensions[gl(x)].width = 3

    # 4) repoint formulas outside the area that referenced moved interim columns
    repointed = []
    if moved:
        pat = re.compile(r"(?<![A-Z$!])(\$?)(" + "|".join(moved) + r")(\$?\d+)\b")
        for row in ws.iter_rows():
            for cell in row:
                if cell.column_letter in area_cols or not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                new = pat.sub(lambda m: m.group(1) + moved[m.group(2)] + m.group(3), cell.value)
                if new != cell.value:
                    repointed.append((cell.coordinate, cell.value, new))
                    cell.value = new

    # 5) reconciliation
    ev = Evaluator(ws)
    keys = [("rev", "Revenue"), ("itax", "Income tax"), ("ni", "NI attrib")]
    pbt = next((r for r in formula_rows if str(ws.cell(r, 2).value).lower().startswith("profit before")), None)
    print("rows:", rows, "| formula rows:", formula_rows, "| yoy rows:", yoy_rows)
    print("period".ljust(8), *[k[1].rjust(14) for k in keys], "PBT".rjust(14))
    for lab, col in pos.items():
        vals = [ev(f"{col}{rows[k]}") for k, _ in keys] + [ev(f"{col}{pbt}") if pbt else None]
        print(lab.ljust(8), *[(f"{v:,.0f}" if isinstance(v, (int, float)) else "-").rjust(14) for v in vals])
    for k, a, parts in cfg.get("checks", []):
        tot = ev(f"{pos[a]}{rows[k]}")
        s = sum(ev(f"{pos[p]}{rows[k]}") or 0 for p in parts)
        print(f"check {k}: {a} {tot:,.1f} vs sum{parts} {s:,.1f} diff {tot - s:,.3f}")
    if repointed:
        print("repointed:", repointed)

    if not dry:
        wb.calculation.fullCalcOnLoad = True
        wb.save(model)
        print("saved", model)
    else:
        print("dry run: not saved")


if __name__ == "__main__":
    main()
