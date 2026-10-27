"""Overwrite a workbook's Guidance sheet with sectioned, numbered guidance sentences.

Usage:
  python write_guidance.py <workbook.xlsx> <guidance.json> [--backup-dir DIR]
                           [--sheet Guidance] [--force]

guidance.json:
  {"title": "...", "sections": [{"heading": "...", "items": ["...", "..."]}]}

Safety steps:
  - aborts if the workbook is open in Excel (~$ lock file next to it)
  - aborts if the workbook has charts/images (openpyxl drops them) unless --force
  - copies the workbook to --backup-dir (default: system temp) before writing
  - after saving, compares every other sheet with the backup and prints differences
"""
import argparse
import json
import math
import os
import shutil
import sys
import tempfile
import zipfile
from datetime import datetime

import openpyxl
from openpyxl.styles import Font


def has_drawings(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    # .vml drawings hold cell comments, which openpyxl keeps; charts/images it drops.
    return [n for n in names if n.startswith(("xl/charts/", "xl/media/"))
            or (n.startswith("xl/drawings/") and not n.endswith(".vml"))]


def same(a, b):
    if isinstance(a, float) and isinstance(b, (int, float)) or isinstance(b, float) and isinstance(a, (int, float)):
        return math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-9)
    return a == b


def compare(backup, saved, skip):
    a = openpyxl.load_workbook(backup)
    b = openpyxl.load_workbook(saved)
    diffs = []
    for name in a.sheetnames:
        if name == skip:
            continue
        if name not in b.sheetnames:
            diffs.append(f"{name}: sheet missing after save")
            continue
        wa, wb_ = a[name], b[name]
        rows = max(wa.max_row, wb_.max_row)
        cols = max(wa.max_column, wb_.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                ca, cb = wa.cell(r, c), wb_.cell(r, c)
                if not same(ca.value, cb.value):
                    diffs.append(f"{name}!{ca.coordinate}: {ca.value!r} -> {cb.value!r}")
                if (ca.comment is None) != (cb.comment is None):
                    diffs.append(f"{name}!{ca.coordinate}: comment changed")
    return diffs


def main():
    p = argparse.ArgumentParser()
    p.add_argument("workbook")
    p.add_argument("guidance_json")
    p.add_argument("--backup-dir", default=tempfile.gettempdir())
    p.add_argument("--sheet", default="Guidance")
    p.add_argument("--force", action="store_true")
    args = p.parse_args()

    wb_path = os.path.abspath(args.workbook)
    folder, fname = os.path.split(wb_path)
    if os.path.exists(os.path.join(folder, "~$" + fname)):
        sys.exit(f"ABORT: {fname} looks open in Excel (lock file found). Close it and retry.")

    drawings = has_drawings(wb_path)
    if drawings and not args.force:
        sys.exit("ABORT: workbook contains charts/images that openpyxl would drop on save:\n  "
                 + "\n  ".join(drawings) + "\nRe-run with --force only if the user accepts that.")

    with open(args.guidance_json, encoding="utf-8") as f:
        data = json.load(f)

    os.makedirs(args.backup_dir, exist_ok=True)
    stem, ext = os.path.splitext(fname)
    backup = os.path.join(args.backup_dir, f"{stem}.backup-{datetime.now():%Y%m%d-%H%M%S}{ext}")
    shutil.copy2(wb_path, backup)

    wb = openpyxl.load_workbook(wb_path)
    if args.sheet in wb.sheetnames:
        ws = wb[args.sheet]
        old = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        for row in ws.iter_rows():
            for c in row:
                c.value = None
    else:
        ws = wb.create_sheet(args.sheet)
        old = []

    r, n = 2, 1
    if data.get("title"):
        ws.cell(r, 2, data["title"]).font = Font(name="Arial", size=11, bold=True)
        r += 2
    for sec in data["sections"]:
        ws.cell(r, 2, sec["heading"]).font = Font(name="Arial", size=10, bold=True)
        r += 1
        for item in sec["items"]:
            ws.cell(r, 2, f"{n}. {item}").font = Font(name="Arial", size=10)
            n += 1
            r += 1
        r += 1
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 110)
    wb.save(wb_path)

    diffs = compare(backup, wb_path, args.sheet)
    print(f"Wrote {n - 1} items in {len(data['sections'])} sections to '{args.sheet}'.")
    print(f"Replaced {len(old)} non-empty cells that were there before.")
    print(f"Backup: {backup}")
    if diffs:
        print(f"WARNING: {len(diffs)} differences in other sheets:")
        for d in diffs[:50]:
            print("  " + d)
    else:
        print("Other sheets: unchanged.")


if __name__ == "__main__":
    main()
