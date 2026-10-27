"""Print every non-empty cell of a workbook, one row per line, with computed values.

Usage: python dump_workbook.py <workbook.xlsx> [--formulas]

By default prints cached values (what Excel last calculated). With --formulas,
prints formulas instead. Floats are rounded to 3 decimals to keep the dump short.
Set PYTHONIOENCODING=utf-8 on Windows so Chinese text prints.
"""
import sys

import openpyxl


def fmt(v):
    if isinstance(v, float):
        return str(round(v, 3))
    return str(v)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    formulas = "--formulas" in sys.argv[2:]
    wb = openpyxl.load_workbook(path, data_only=not formulas)
    for ws in wb:
        print(f"===== {ws.title} ({ws.max_row} rows x {ws.max_column} cols)")
        for row in ws.iter_rows():
            cells = [f"{c.coordinate}={fmt(c.value)}" for c in row if c.value is not None]
            if cells:
                print(" | ".join(cells))


if __name__ == "__main__":
    main()
