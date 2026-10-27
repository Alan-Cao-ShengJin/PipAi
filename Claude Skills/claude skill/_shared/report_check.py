# -*- coding: utf-8 -*-
"""报告硬性规范自检器 —— 把各 SKILL.md 里的 HARD RULES 变成会失败的断言。

以前这些规则只写在 SKILL.md 里靠人记，结果每次都在交付前才发现违规
（表格论证不足、段落超长、字面 `**` 漏进正文、898 个空 run 污染字体校验）。
现在统一在这里检查。

两个入口：

    from report_check import check_blocks, check_docx
    bad = check_blocks(r.blocks, core_tables={'表2','表3','表5','表6','表17'})
    bad += check_docx('out/x.docx')
    if bad: raise SystemExit('\\n'.join(bad))

规则来源：
  penetrate-narrative-stock-analysis SKILL.md 「篇幅与摘要要求」「表格论证规范」
  penetrate-financial-report-v2      SKILL.md 「五、格式与样式规范」「七、执行检查清单」
"""
import re

from docx import Document
from docx.oxml.ns import qn

# --- UTF-8 stdout (Windows cp1252 consoles mangle CJK without this) ---
import sys as _sys
try:
    _sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    _sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

FONT = '微软雅黑'
PARA_HARD = 330          # 单段字符上限（≈6 行的宽松判定）
ABSTRACT_MIN = 800
CONCLUSION_EACH = 150
CONCLUSION_TOTAL = 900
TABLE_GENERAL = 300
TABLE_CORE = 500

# 交付物里绝不允许出现的来源标记（narrative skill v3.0.0）
SOURCE_PAT = re.compile(
    r'\[source:\s*\d+\]'
    r'|\[来源[:：][^\]]*\]'
    r'|\[(?:report|roadShow|comment|edb)_[A-Za-z0-9_\-]+\]'
)
# 既不是 Markdown 也不是 HTML 的渲染器，字面 ** 会原样印出来
MD_PAT = re.compile(r'\*\*|^\s*#{1,6}\s|^\s*[-*]\s+\S', re.M)


def _cn(s):
    """字数 = 非空白字符数。

    口径说明（这里踩过一次）：SKILL.md 写的是「中文字符计」，最初实现成只数汉字。
    但财报/估值类正文里「49.58 亿元」「L/E3 = 1.17 倍」这类数字与符号本身就是论证，
    只数汉字会把实打实的内容判成字数不足——用已交付的 002749 报告回归时，
    9 张合规的表被误报为论证不足。所以改为数非空白字符，接近中文语境下「字数」的常识。
    """
    return len(re.sub(r'\s', '', s or ''))


def check_blocks(blocks, core_tables=(), abstract_head='摘要',
                 conclusion_head='核心结论', expect_conclusions=6,
                 require_lengths=True):
    """渲染前自检。返回问题字符串列表，空列表＝全过。

    core_tables: 需要 500 字论证的核心表编号集合，如 {'表2','表3','表5'}。
                 匹配方式是 cap 文本里出现该子串。
    """
    bad = []
    texts = [(k, v) for k, v in blocks if k in ('p', 'quote', 'cap', 'note',
                                                'h1', 'h2', 'h3')]

    # 1) 段落长度 & 空段落
    for k, v in texts:
        if k in ('p', 'quote'):
            if not (v or '').strip():
                bad.append(f'[空段落] 出现空的 {k} block')
            elif len(v) > PARA_HARD:
                bad.append(f'[段落超长] {len(v)} 字 > {PARA_HARD}：{v[:40]}…')

    # 2) 来源标记 / 字面 Markdown
    for k, v in texts:
        if isinstance(v, str):
            m = SOURCE_PAT.search(v)
            if m:
                bad.append(f'[来源标记] 交付物不得出现 "{m.group(0)}"：{v[:40]}…')
            m = MD_PAT.search(v)
            if m:
                bad.append(f'[字面Markdown] 渲染器不处理 Markdown，'
                           f'"{m.group(0).strip()}" 会原样印出：{v[:40]}…')
    for k, v in blocks:
        if k == 'table':
            for row in v[1]:
                for cell in row:
                    if isinstance(cell, str) and SOURCE_PAT.search(cell):
                        bad.append(f'[来源标记] 表格单元格：{cell[:40]}…')

    if not require_lengths:
        return bad

    # 3) 摘要字数
    abs_txt = _section_text(blocks, abstract_head)
    if abs_txt is None:
        bad.append(f'[缺章节] 找不到「{abstract_head}」章节')
    elif _cn(abs_txt) < ABSTRACT_MIN:
        bad.append(f'[摘要过短] {_cn(abs_txt)} 字 < {ABSTRACT_MIN}')

    # 4) 核心结论：条数与每条字数
    #    一条结论 = 一个 h2/h3 小标题 + 其下若干段（判断/依据/数据/推论）。
    #    「每段≤6 行」这条规则本身就要求把一条结论拆成多段，所以绝不能
    #    「一段 = 一条结论」地去数——那样每条都会被误判为过短。
    cons = _conclusions(blocks, conclusion_head)
    if cons is None:
        bad.append(f'[缺章节] 找不到「{conclusion_head}」章节')
    else:
        if len(cons) < expect_conclusions:
            bad.append(f'[结论条数] 只有 {len(cons)} 条实质结论，'
                       f'要求 {expect_conclusions} 条')
        for i, (name, n) in enumerate(cons[:expect_conclusions], 1):
            if n < CONCLUSION_EACH:
                bad.append(f'[结论过短] 第{i}条「{name[:16]}」{n} 字 < {CONCLUSION_EACH}')
        tot = sum(n for _, n in cons)
        if tot < CONCLUSION_TOTAL:
            bad.append(f'[结论合计过短] {tot} 字 < {CONCLUSION_TOTAL}')

    # 5) 一表一论：每个 cap 之后、下一个 cap 之前的正文字数
    #    引用块（表4/7/9/11/15 那种「关键判断引用块」）同样算一张表，
    #    它们也有 300 字论证下限，不能因为没有 <table> 就漏检。
    for cap, words, tbls in _table_sections(blocks):
        if tbls == 0:
            continue
        core = any(c in cap for c in core_tables)
        floor = TABLE_CORE if core else TABLE_GENERAL
        if words < floor:
            bad.append(f'[表格论证不足] 「{cap[:28]}」{words} 字 < {floor}'
                       f'{"（核心表）" if core else ""}')
        if tbls > 1:
            bad.append(f'[多表共用论证] 「{cap[:28]}」下挂了 {tbls} 张表，'
                       f'规范要求一表一论')
    return bad


def _conclusions(blocks, head):
    """返回 [(小标题, 字数), ...]。一条结论 = 一个小标题 + 其下所有正文段。

    没有小标题时退回「一段一条」（少数报告把六条写成六个大段）。
    """
    on, lvl, out, cur = False, 9, [], None
    loose = []
    for k, v in blocks:
        if k in ('h1', 'h2', 'h3'):
            cur_lvl = int(k[1])
            if on and cur_lvl <= lvl:
                break                      # 离开「核心结论」章节
            if not on and isinstance(v, str) and head in v:
                on, lvl = True, cur_lvl
                continue
            if on:
                out.append([v, 0])         # 一条新结论开始
            continue
        if on and k in ('p', 'quote'):
            if out:
                out[-1][1] += _cn(v)
            else:
                loose.append(v)            # 小标题之前的散段
    if not on:
        return None
    if out:
        return [(n, c) for n, c in out]
    return [(p[:16], _cn(p)) for p in loose if _cn(p) >= 40]


def _section_text(blocks, head):
    ps = _section_paras(blocks, head)
    return None if ps is None else ''.join(ps)


def _section_paras(blocks, head):
    """取某个标题之下、到下一个同级或更高级标题之前的所有正文段落。"""
    out, on, lvl = [], False, 9
    for k, v in blocks:
        if k in ('h1', 'h2', 'h3'):
            cur = int(k[1])
            if on and cur <= lvl:
                break
            if not on and isinstance(v, str) and head in v:
                on, lvl = True, cur
            continue
        if on and k in ('p', 'quote'):
            out.append(v)
    return out if on else None


def _table_sections(blocks):
    """把 blocks 切成 (表题, 该表论证字数, 该段内表格数) 三元组。

    引用块也计入表格数：SKILL.md 的 17 张表里，表4/7/9/11/15 本来就是
    「关键判断引用块」，同样有论证字数下限。引用块自身的文字不计入论证字数
    （它是被论证的对象，不是论证）。
    """
    secs, cap, words, tbls = [], None, 0, 0
    for k, v in blocks:
        if k == 'cap':
            if cap is not None:
                secs.append((cap, words, tbls))
            cap, words, tbls = v, 0, 0
        elif cap is not None:
            if k == 'table':
                tbls += 1
            elif k == 'quote':
                tbls += 1
            elif k == 'p':
                words += _cn(v)
            elif k == 'h1':
                # 只在「章」的边界断开。表格的论证常常写在紧随其后的 h2 小节里
                # （如表6B 之后的「逐家终局预期反算」），在 h2 处断开会把那段
                # 论证判成 0 字 —— 这是回归 002749 时暴露出来的误报。
                secs.append((cap, words, tbls))
                cap, words, tbls = None, 0, 0
    if cap is not None:
        secs.append((cap, words, tbls))
    return secs


def check_docx(path, body_pt=10.0, table_pt=9.0):
    """存盘后复核：字体、字号、段间距、空段落。逐 run 检查，不抽样。"""
    bad = []
    doc = Document(path)
    tbl_runs = set()
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        tbl_runs.add(id(r._element))
                        _chk_run(r, table_pt, bad, '表格')

    empties = 0
    for p in doc.paragraphs:
        if not p.text.strip():
            # 只放了图片的段落不是空段落 —— 别把 <w:drawing> 误判成空行
            if p._element.findall('.//' + qn('w:drawing')):
                continue
            empties += 1          # 无文字也无图 = 空段落（含零 run 的那种）
            continue
        pf = p.paragraph_format
        for r in p.runs:
            if id(r._element) in tbl_runs:
                continue
            nm = r.font.name
            if nm != FONT:
                bad.append(f'[字体] 非{FONT}（{nm}）："{r.text[:26]}"')
            rPr = r._element.find(qn('w:rPr'))
            rf = rPr.find(qn('w:rFonts')) if rPr is not None else None
            if rf is None or rf.get(qn('w:eastAsia')) != FONT:
                bad.append(f'[eastAsia未绑定] 中文会回退宋体："{r.text[:26]}"')
        # 3磅间距只约束「正文段落」。标题、表题(9.5pt)、附注(8pt)、图注(8.5pt)、
        # 副标题(10.5pt) 是辅助文字，各有自己的呼吸量，不在此规则内。
        sz = {round(r.font.size.pt, 2) for r in p.runs if r.font.size}
        if sz and sz == {round(body_pt, 2)}:
            for name, val in (('段前', pf.space_before), ('段后', pf.space_after)):
                got = 'None' if val is None else f'{val.pt:g}磅'
                if val is None or abs(val.pt - 3) > 0.01:
                    bad.append(f'[{name}间距] 应 3磅，实为 {got}："{p.text[:26]}"')
    if empties:
        bad.append(f'[空段落] {empties} 个空段落，规范要求靠 3磅间距区隔而非空行')
    return bad


def _chk_run(r, pt, bad, where):
    if not r.text:
        return
    if r.font.name != FONT:
        bad.append(f'[{where}字体] 非{FONT}（{r.font.name}）："{r.text[:20]}"')
    if r.font.size and abs(r.font.size.pt - pt) > 0.01:
        bad.append(f'[{where}字号] 应 {pt}pt，实为 {r.font.size.pt}pt'
                   f'："{r.text[:20]}"')


def report(bad, label='报告自检'):
    """打印结果并返回是否通过。"""
    if not bad:
        print(f'{label}：全部通过 ✓')
        return True
    print(f'{label}：{len(bad)} 项不合规')
    for b in bad:
        print('  ✗', b)
    return False
