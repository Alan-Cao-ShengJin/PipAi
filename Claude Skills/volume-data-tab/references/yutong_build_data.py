import os, shutil, sys, warnings
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

warnings.filterwarnings("ignore")
D = r"C:\Users\Alan Cao\Desktop\Companies\宇通"
SCR = os.path.dirname(os.path.abspath(__file__))
MODEL = os.path.join(D, "宇通客车 1015.xlsx")
GX = os.path.join(D, "宇通客车拆分【国信汽车】261014.xlsx")
TR = os.path.join(D, "宇通月销跟踪(更新至7月）.xlsx")

if os.path.exists(os.path.join(D, "~$宇通客车 1015.xlsx")):
    sys.exit("model open in Excel")
backup = os.path.join(SCR, f"宇通客车 1015.backup-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
shutil.copy2(MODEL, backup)

gx = openpyxl.load_workbook(GX, data_only=True)
tr = openpyxl.load_workbook(TR, data_only=True)
SRC_MODEL = os.path.join(SCR, "宇通客车 1015.backup-20261015-141116.xlsx")  # last version with Monthly tab
mdl_v = openpyxl.load_workbook(SRC_MODEL, data_only=True)

# ---------------- monthly data ----------------
months = [f"{y}-{m:02d}" for y in range(2022, 2027) for m in range(1, 13)]
months = [m for m in months if m <= "2026-09"]
rec = {m: {} for m in months}
note = {m: [] for m in months}
derived = set()  # (month, field)

# 2022-01..2024-11 production & sales (total / 大中 / 轻) from tracker Sheet4
s4 = tr["Sheet4"]
for r in range(2, 40):
    a = s4.cell(r, 1).value
    if isinstance(a, datetime):
        m = f"{a.year}-{a.month:02d}"
        v = [s4.cell(r, c).value for c in range(2, 9)]
        rec[m].update(prod=v[0], sales=v[1], prod_dz=v[2], sales_dz=v[3], prod_q=v[5], sales_q=v[6])

# 2023-01..2026-09 sales split from 国信 月度跟踪 (rows 4..48, sequential months)
mt = gx["数据月度跟踪-更新至9月"]
for i, r in enumerate(range(4, 49)):
    m = months[12 + i]
    g = lambda c: mt.cell(r, c).value
    rec[m].update(exp=g(4), nev=g(5), dom=g(6), dom_dz=g(7), sales_dz=g(8), sales_q=g(9), sales=g(10))
# Jun-2026: 国信 has 5231/4308/2108; company announcement (国信 Sheet4, tracker, Monthly tab) = 5213 / 4290
rec["2026-06"].update(sales=5213, sales_dz=4290, dom_dz=2090)
note["2026-06"].append("销量按公司公告5213（国信月度表为5231）")
note["2026-07"].append("出口/新能源取国信1182/330（月销跟踪为1200/300）")
rec["2023-04"]["exp"] = 976
note["2023-04"].append("出口按 大中2497−国内大中1521 修正为976（两份来源均为972，与合计2836差4辆）")
for m in months:
    e, t, d0 = rec[m].get("exp"), rec[m].get("sales"), rec[m].get("dom")
    if e is not None and d0 is not None and t - e != d0:
        note[m].append(f"国内=合计−出口={t-e}（来源原值{d0}）")

# 2025/2026 Jan-Jul by length from Monthly tab
mo = mdl_v["Monthly"]
for yr, c0 in ((2025, 2), (2026, 11)):
    for k in range(7):
        m = f"{yr}-{k+1:02d}"
        c = c0 + k
        pb, pm, pq = (mo.cell(r, c).value for r in (3, 4, 5))
        sb, sm, sq = (mo.cell(r, c).value for r in (9, 10, 11))
        rec[m].update(prod_b=pb, prod_m=pm, prod_q=pq, prod_dz=pb + pm, prod=pb + pm + pq,
                      sales_b=sb, sales_m=sm)
        ind = mo.cell(15, c).value
        if ind is not None:
            rec[m]["ind_exp"] = ind

# Sep 2025 / Sep 2026 + YTD from 国信 Sheet4 (company announcement); Aug = YTD Sep - Jan..Jul - Sep
an = gx["Sheet4"]
def col(cn):  # rows 5..12 : prod, 大, 中, 轻, sales, 大, 中, 轻
    return [an.cell(r, cn).value for r in range(5, 13)]
sep26, sep25, ytd26, ytd25 = col(6), col(7), col(9), col(10)
for yr, sep, ytd in ((2026, sep26, ytd26), (2025, sep25, ytd25)):
    s = f"{yr}-09"
    rec[s].update(prod=sep[0], prod_b=sep[1], prod_m=sep[2], prod_q=sep[3], prod_dz=sep[1] + sep[2],
                  sales_b=sep[5], sales_m=sep[6])
    a = f"{yr}-08"
    j7 = [f"{yr}-{k:02d}" for k in range(1, 8)]
    pb = ytd[1] - sum(rec[m]["prod_b"] for m in j7) - sep[1]
    pm = ytd[2] - sum(rec[m]["prod_m"] for m in j7) - sep[2]
    pq = ytd[3] - sum(rec[m]["prod_q"] for m in j7) - sep[3]
    sb = ytd[5] - sum(rec[m]["sales_b"] for m in j7) - sep[5]
    sm = ytd[6] - sum(rec[m]["sales_m"] for m in j7) - sep[6]
    rec[a].update(prod_b=pb, prod_m=pm, prod_q=pq, prod_dz=pb + pm, prod=pb + pm + pq, sales_b=sb, sales_m=sm)
    for f in ("prod_b", "prod_m", "prod_q", "prod_dz", "prod", "sales_b", "sales_m"):
        derived.add((a, f))
    note[a].append("产量及大/中型销量=9月累计−1~7月−9月（推算）")
    # sanity: derived 大+中+轻 should equal reported monthly total
    assert sb + sm + rec[a]["sales_q"] == rec[a]["sales"], (a, sb, sm, rec[a]["sales_q"], rec[a]["sales"])


# ---------------- FactSet StreetAccount (company monthly announcements) ----------------
FS_SALES = {"2022-01":1828,"2022-02":840,"2022-03":2277,"2022-04":2151,"2022-05":2237,"2022-06":2234,"2022-07":2040,"2022-08":2842,
 "2022-09":2510,"2022-10":2884,"2022-11":3586,"2022-12":4769,"2023-01":1261,"2023-02":788,"2023-03":2390,"2023-04":2836,"2023-05":3133,
 "2023-06":4726,"2023-07":3675,"2023-08":3899,"2023-09":3567,"2023-10":2249,"2023-11":3285,"2023-12":4709,"2024-01":2797,"2024-02":1385,
 "2024-03":3549,"2024-04":4598,"2024-05":3483,"2024-06":4743,"2024-07":3355,"2024-08":3648,"2024-09":3788,"2024-10":3221,"2024-11":3736,
 "2024-12":8615,"2025-01":2550,"2025-02":1571,"2025-03":4890,"2025-04":3338,"2025-05":3053,"2025-06":5919,"2025-07":3219,"2025-08":4260,
 "2025-09":4756,"2025-10":3040,"2025-11":4058,"2025-12":8864,"2026-01":1728,"2026-02":1806,"2026-03":4118,"2026-04":3755,"2026-05":3480,
 "2026-06":5213,"2026-07":2925,"2026-08":3746,"2026-09":4158}
FS_PROD = {"2023-02":1206,"2023-03":2566,"2023-04":3093,"2023-05":3616,"2023-06":4093,"2023-07":3778,"2023-08":4056,"2023-09":3407,
 "2023-10":2577,"2023-11":3342,"2023-12":4315,"2024-01":2997,"2024-02":1325,"2024-03":4039,"2024-04":4977,"2024-05":3665,"2024-06":4257,
 "2024-07":2600,"2024-08":4189,"2024-09":3878,"2024-10":3558,"2024-11":4665,"2024-12":6904,"2025-01":2560,"2025-02":1578,"2025-03":4796,
 "2025-04":3517,"2025-05":3434,"2025-06":5528,"2025-07":3849,"2025-08":4165,"2025-09":4447,"2025-10":3228,"2025-11":5188,"2025-12":7066,
 "2026-01":2167,"2026-02":2061,"2026-03":4201,"2026-04":4204,"2026-05":3570,"2026-06":4500,"2026-07":3191,"2026-08":3681,"2026-09":4149}
note["2024-01"].append("FactSet 2024-01快讯把产销写反（销2997/产2797），按公司公告及FactSet 2025-01同比基数取销2797/产2997")
fs_filled = set()
for m, v in FS_PROD.items():
    if rec[m].get("prod") is None:
        rec[m]["prod"] = v
        fs_filled.add((m, "prod"))
        note[m].append("产量取自FactSet")
mism = [(m, f, rec[m].get(f), src[m]) for f, src in (("sales", FS_SALES), ("prod", FS_PROD)) for m in src if rec[m].get(f) != src[m]]
print("mismatch vs FactSet:", mism)

# ---------------- write ----------------
wb = openpyxl.load_workbook(MODEL)
if "Data" in wb.sheetnames:
    del wb["Data"]
ws = wb.create_sheet("Data")
BLUE = Font(name="Arial", size=9, color="0000CC")
BLACK = Font(name="Arial", size=9)
ORANGE = Font(name="Arial", size=9, color="C65911", italic=True)
GREY = Font(name="Arial", size=9, color="7F7F7F")
BOLD = Font(name="Arial", size=9, bold=True)
TITLE = Font(name="Arial", size=11, bold=True)
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
INT, PCT, YI = "#,##0;-#,##0", "0.0%", "#,##0.00"

def title(r, text, src):
    ws.cell(r, 2, text).font = TITLE
    ws.cell(r + 1, 2, src).font = GREY
    return r + 2

def header(r, labels, c0=2):
    for i, h in enumerate(labels):
        c = ws.cell(r, c0 + i, h)
        c.font, c.fill = BOLD, HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30

def put(r, c, v, fmt=INT, font=BLUE):
    if v is None:
        return
    cell = ws.cell(r, c, v)
    cell.font = font if not (isinstance(v, str) and v.startswith("=")) else BLACK
    cell.number_format = fmt

ws.cell(1, 2, "宇通客车 数据汇总（Data）").font = Font(name="Arial", size=13, bold=True)
ws.cell(2, 2, "字体颜色：蓝=来源数据；黑=公式；橙色斜体=推算值；灰=第三方估算/预测。单位见各表。").font = GREY

# ---- Block 1 monthly ----
r = title(4, "一、月度产销（辆）",
          "来源：2022–2024产量&大中/轻客销量=宇通月销跟踪 Sheet4（公司公告）；2023-01起内外销拆分=国信 数据月度跟踪-更新至9月；"
          "2025/2026年1–7月大/中/轻型=原Monthly tab；9月=公司公告（国信 Sheet4）；行业出口=原Monthly tab；FactSet=StreetAccount月度产销快讯（公司公告）。2024-12、2025-10~12产量取自FactSet。出口/出口新能源/国内大中为第三方口径；国内=合计−出口（公式），保证国内+出口=合计。")
mcols = ["月份", "产量合计", "产量-大型", "产量-中型", "产量-大中", "产量-轻型", "销量合计", "销量-大型", "销量-中型",
         "销量-大中", "销量-轻型", "国内", "国内大中", "出口", "出口新能源", "出口传统", "出口占比", "出口新能源占出口", "行业出口(大中型)", "FactSet销量", "FactSet产量", "销量差异(本表−FactSet)", "备注"]
fields = [None, "prod", "prod_b", "prod_m", "prod_dz", "prod_q", "sales", "sales_b", "sales_m", "sales_dz", "sales_q",
          "dom", "dom_dz", "exp", "nev", None, None, None, "ind_exp", None, None, None, None]
header(r, mcols)
mrow = {}
r += 1
for m in months:
    mrow[m] = r
    ws.cell(r, 2, m).font = BOLD
    for i, f in enumerate(fields):
        if f is None:
            continue
        v = rec[m].get(f)
        put(r, 2 + i, v, INT, ORANGE if (m, f) in derived else BLUE)
    if rec[m].get("exp") is not None:
        put(r, 13, f"=H{r}-O{r}")
        put(r, 17, f"=O{r}-P{r}")
        put(r, 18, f"=IFERROR(O{r}/H{r},\"\")", PCT)
        put(r, 19, f"=IFERROR(P{r}/O{r},\"\")", PCT)
    put(r, 21, FS_SALES.get(m))
    put(r, 22, FS_PROD.get(m))
    put(r, 23, f'=IF(U{r}="","",H{r}-U{r})')
    if note[m]:
        ws.cell(r, 24, "；".join(note[m])).font = GREY
    r += 1
m_last = r - 1

# ---- Block 2 quarterly ----
r += 2
r = title(r, "二、季度销量（辆）",
          "来源：2023Q1起=上表月度加总（公式）；2021Q4–2022Q4=国信 数据月度跟踪。原Quarterly tab的1Q25/1Q26数（出口2068/2689、新能源275/440）与月度加总（2050/2660、290/466）不同，以月度加总为准。")
header(r, ["季度", "销量合计", "大中", "轻型", "国内", "国内大中", "出口", "出口传统", "出口新能源", "出口占比", "出口新能源占出口"])
early = {"2021Q4": 155, "2022Q1": 156, "2022Q2": 157, "2022Q3": 158, "2022Q4": 159}
qtrs = list(early) + [f"{y}Q{n}" for y in range(2023, 2027) for n in range(1, 5) if f"{y}Q{n}" <= "2026Q3"]
r += 1
for q in qtrs:
    ws.cell(r, 2, q).font = BOLD
    y, n = int(q[:4]), int(q[-1])
    ms = [f"{y}-{3*(n-1)+k:02d}" for k in (1, 2, 3)]
    if q in early:
        er = early[q]
        g = lambda c: mt.cell(er, c).value
        for c, v in {3: g(10), 4: g(8), 5: g(9), 6: g(6), 7: g(7), 8: g(4), 10: g(5)}.items():
            put(r, c, v)
    else:
        a_, b_ = mrow[ms[0]], mrow[ms[2]]
        for c, mc in ((3, "H"), (4, "K"), (5, "L"), (6, "M"), (7, "N"), (8, "O"), (10, "P")):
            put(r, c, f"=SUM({mc}{a_}:{mc}{b_})")
    put(r, 9, f"=H{r}-J{r}")
    put(r, 11, f"=IFERROR(H{r}/C{r},\"\")", PCT)
    put(r, 12, f"=IFERROR(J{r}/H{r},\"\")", PCT)
    r += 1

# ---- Block 2b annual reconciliation ----
r += 2
r = title(r, "二b、年度汇总与年报核对（辆）",
          "月度/季度加总 vs 公司年报口径（年报数取自 Revenue Projection）。出口及新能源差异来自第三方月度口径与公司年报口径不同；销量合计应为0差异。")
header(r, ["年份", "销量合计", "年报销量", "差异", "出口", "年报出口", "差异", "出口新能源", "年报出口新能源(纯电)", "差异", "国内", "年报国内", "差异"])
rp = mdl_v["Revenue Projection"]
r += 1
qrow = {}
rr0 = r - 1
# locate quarterly rows by label
for rr in range(1, r):
    v = ws.cell(rr, 2).value
    if isinstance(v, str) and len(v) == 6 and v[4] == "Q":
        qrow[v] = rr
for yr, col in ((2022, 7), (2023, 8), (2024, 9), (2025, 10), (2026, None)):
    qs = [qrow[f"{yr}Q{n}"] for n in range(1, 5) if f"{yr}Q{n}" in qrow]
    ws.cell(r, 2, str(yr) if yr < 2026 else "2026年1–9月").font = BOLD
    def sumq(c):
        return "=" + "+".join(f"{L(c)}{x}" for x in qs)
    for c_out, c_q, ar_row in ((3, 3, 68), (6, 8, 65), (9, 10, 78), (12, 6, 62)):
        put(r, c_out, sumq(c_q))
        if col:
            put(r, c_out + 1, f"='Revenue Projection'!{L(col)}{ar_row}")
            put(r, c_out + 2, f"={L(c_out)}{r}-{L(c_out+1)}{r}")
    r += 1

# ---- Block 3 annual exports by region ----
r += 2
r = title(r, "三、年度出口分区域（行业：万台；宇通：万台/台）",
          "来源：国信 年度出口数据 拆分（海外、公司）；行业=公司官方PPT，宇通=海关数据")
rg = gx["年度出口数据 拆分（海外、公司）"]
blocks = [(4, 11, "行业海外销量（万台）"), (13, 20, "行业新能源海外销量（万台）"), (22, 29, "新能源渗透率"),
          (31, 39, "宇通出口销量（万台，海关）"), (41, 49, "宇通出口电车销量（台，海关）")]
region_notes = []
header(r, ["项目", "2022", "2023", "2024", "2025"])
r += 1
for a, b, name in blocks:
    for rr in range(a, b + 1):
        lab = rg.cell(rr, 2).value
        if lab is None:
            continue
        ws.cell(r, 2, name if rr == a else "  " + str(lab)).font = BOLD if rr == a else BLACK
        fmt = PCT if "渗透率" in name else (YI if "万台" in name else INT)
        for k in range(4):
            if rr == a and "渗透率" not in name:
                cl = L(3 + k)
                put(r, 3 + k, f"=SUM({cl}{r+1}:{cl}{r+b-a})", fmt)
                orig = rg.cell(rr, 3 + k).value
                sub = sum(rg.cell(x, 3 + k).value or 0 for x in range(a + 1, b + 1))
                if orig is not None and abs(orig - sub) > 1e-9:
                    region_notes.append(f"{name} {2022+k}: 原合计{orig}，分区域加总{round(sub,4)}")
            else:
                put(r, 3 + k, rg.cell(rr, 3 + k).value, fmt)
        r += 1

if region_notes:
    ws.cell(r, 2, "合计改为分区域SUM公式；原表合计因四舍五入有差：" + "；".join(region_notes)).font = GREY
    r += 1

# ---- Block 4 customs NEV exports 2026 ----
r += 2
r = title(r, "四、2026年新能源客车出口海关数据（台）",
          "来源：国信 海关数据更新至7月（实际为1–6月）；公司口径出口新能源=国信 Sheet2（开票口径）。海关为报关口径，与公司交付口径有时间差。")
cu = gx["海关数据更新至7月"]
cust, cur = {}, None
for rr in range(4, 107):
    a, b, c = (cu.cell(rr, k).value for k in (1, 2, 3))
    if a is None:
        continue
    a = str(a).strip()
    if a.isdigit() and len(a) == 6:
        cur = a
        continue
    cust.setdefault(a, {})[cur] = (b or 0, c or 0)
mons = ["202601", "202602", "202603", "202604", "202605", "202606"]
header(r, ["国家/月份"] + [f"{m[:4]}-{m[4:]}" for m in mons] + ["1H合计(台)"])
r += 1
cfirst = r
for ctry in sorted(cust, key=lambda k: -sum(v[0] for v in cust[k].values())):
    ws.cell(r, 2, ctry).font = BLACK
    for i, m in enumerate(mons):
        if m in cust[ctry]:
            put(r, 3 + i, cust[ctry][m][0])
    put(r, 9, f"=SUM(C{r}:H{r})")
    r += 1
clast = r - 1
ws.cell(r, 2, "海关合计").font = BOLD
for i in range(6):
    cl = L(3 + i)
    put(r, 3 + i, f"=SUM({cl}{cfirst}:{cl}{clast})")
put(r, 9, f"=SUM(I{cfirst}:I{clast})")
tot_r = r
r += 1
s2 = gx["Sheet2"]
ws.cell(r, 2, "公司口径出口新能源").font = BOLD
for i in range(6):
    put(r, 3 + i, s2.cell(12 + i, 6).value)
put(r, 9, f"=SUM(C{r}:H{r})")
r += 1

# ---- Block 5 国信 forecasts ----
r += 2
r = title(r, "五、国信销量预测（参考，2026-10-14）", "来源：国信 业绩测算（量利）。非公司指引。")
lv = gx["业绩测算（量利）"]
header(r, ["项目", "2026Q4E", "2026E", "2027E", "2028E"])
r += 1
for row, lab, fmt in ((3, "销量合计（辆）", INT), (6, "国内（辆）", INT), (9, "出口（辆）", INT), (10, "出口传统（辆）", INT),
                      (11, "出口新能源（辆）", INT), (14, "其中欧洲（辆）", INT)):
    ws.cell(r, 2, lab).font = BLACK
    for i, c in enumerate((19, 28, 29, 30)):
        v = lv.cell(row, c).value
        if isinstance(v, (int, float)):
            put(r, 3 + i, v, fmt, GREY)
    r += 1

# layout
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 22
for c in range(3, 25):
    ws.column_dimensions[L(c)].width = 11
ws.column_dimensions["U"].width = 11
ws.freeze_panes = "C4"
ws.column_dimensions["X"].width = 60
for old in ("Monthly", "Quarterly"):
    if old in wb.sheetnames:
        del wb[old]
wb.move_sheet("Data", offset=wb.sheetnames.index("Revenue Projection") + 1 - wb.sheetnames.index("Data"))
wb.calculation.fullCalcOnLoad = True
wb.save(MODEL)
print("backup:", backup)
print("monthly rows", mrow[months[0]], m_last, "| sheets:", wb.sheetnames)
